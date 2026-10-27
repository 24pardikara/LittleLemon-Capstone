"""
Load LittleLemon_data.xlsx into the normalized LittleLemonDB tables.

Reads the Orders sheet, de-duplicates Customers and Menu, inserts Orders,
and seeds a few sample Bookings for the stored procedures.

Run:  py -3.12 load_data.py
"""
import os
import mysql.connector
from openpyxl import load_workbook

DB = dict(host="localhost", user="root",
          password=os.getenv("MYSQL_PASSWORD", "your_password"),  # set MYSQL_PASSWORD env var
          database="LittleLemonDB")
XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LittleLemon_data.xlsx")


def d(v):
    """datetime -> date (or None)."""
    return v.date() if hasattr(v, "date") else v


def main():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Orders"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)  # skip header

    cnx = mysql.connector.connect(**DB)
    cur = cnx.cursor()

    # clear existing data (children first)
    cur.execute("SET FOREIGN_KEY_CHECKS=0")
    for t in ("Orders", "Bookings", "Menu", "Customers"):
        cur.execute(f"TRUNCATE TABLE {t}")
    cur.execute("SET FOREIGN_KEY_CHECKS=1")

    customers = {}          # CustomerID -> (name, city, country, postal, code)
    menus = {}              # composition tuple -> MenuID (assigned later)
    menu_order = []         # keep insertion order
    orders = []             # list of tuples for Orders

    for r in rows:
        (row_no, order_id, order_date, delivery_date, cust_id, cust_name, city,
         country, postal, code, cost, sales, qty, discount, delivery_cost,
         course, cuisine, starter, desert, drink, sides) = r

        if cust_id not in customers:
            customers[cust_id] = (cust_name, city, country, str(postal) if postal is not None else None, code)

        menu_key = (course, cuisine, starter, desert, drink, sides)
        if menu_key not in menus:
            menus[menu_key] = None
            menu_order.append(menu_key)

        orders.append((row_no, order_id, d(order_date), d(delivery_date), cust_id, menu_key,
                       cost, sales, qty, discount, delivery_cost))

    # --- insert Customers ---
    cur.executemany(
        "INSERT INTO Customers (CustomerID, CustomerName, City, Country, PostalCode, CountryCode)"
        " VALUES (%s,%s,%s,%s,%s,%s)",
        [(cid,) + v for cid, v in customers.items()])

    # --- insert Menu (capture generated MenuID for each composition) ---
    for key in menu_order:
        cur.execute(
            "INSERT INTO Menu (CourseName, CuisineName, StarterName, DesertName, Drink, Sides)"
            " VALUES (%s,%s,%s,%s,%s,%s)", key)
        menus[key] = cur.lastrowid

    # --- insert Orders in batches ---
    order_rows = [(rno, oid, od, dd, cid, menus[mk], cost, sales, qty, disc, dcost)
                  for (rno, oid, od, dd, cid, mk, cost, sales, qty, disc, dcost) in orders]
    cur.executemany(
        "INSERT INTO Orders (OrderRecordID, OrderID, OrderDate, DeliveryDate, CustomerID, MenuID,"
        " Cost, Sales, Quantity, Discount, DeliveryCost)"
        " VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", order_rows)

    # --- seed sample Bookings ---
    sample_ids = list(customers.keys())[:3]
    cur.executemany(
        "INSERT INTO Bookings (BookingID, BookingDate, TableNumber, CustomerID) VALUES (%s,%s,%s,%s)",
        [(1, "2022-10-10", 5, sample_ids[0]),
         (2, "2022-11-12", 3, sample_ids[1]),
         (3, "2022-10-11", 2, sample_ids[2])])

    cnx.commit()
    print(f"Loaded: {len(customers)} customers, {len(menu_order)} menu items, "
          f"{len(order_rows)} orders, 3 sample bookings.")
    cur.close()
    cnx.close()


if __name__ == "__main__":
    main()
